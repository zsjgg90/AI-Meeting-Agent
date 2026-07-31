from __future__ import annotations

import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]

REVIEW_FILE = (
    PROJECT_ROOT
    / "data"
    / "eval"
    / "rag_audit"
    / "rag_v2_1"
    / "base_125_review_final_expert_v5.xlsx"
)

OUTPUT_ROOT = PROJECT_ROOT / "data" / "rag" / "v2_1"

POLICY_FILE = OUTPUT_ROOT / "policy" / "policy_fixed_24.jsonl"
DYNAMIC_FILE = OUTPUT_ROOT / "dynamic" / "rag_dynamic_32.jsonl"
SCENARIO_FILE = OUTPUT_ROOT / "scenarios" / "rag_scenario_30.jsonl"

COMBINED_FILE = (
    OUTPUT_ROOT
    / "meeting_analyst_rag_v2_1_86.jsonl"
)

ALIAS_FILE = OUTPUT_ROOT / "aliases" / "alias_trigger_13.json"
MEMORY_FILE = OUTPUT_ROOT / "memory" / "memory_rules_2.jsonl"

DOMAIN_FILE = (
    OUTPUT_ROOT
    / "domain_extensions"
    / "domain_extension_4.json"
)

MERGE_FILE = OUTPUT_ROOT / "merge_map_20.json"
MANIFEST_FILE = OUTPUT_ROOT / "manifest.v2_1.json"


DATASET_VERSION = "meeting_analyst_rag_v2_1"
SCHEMA_VERSION = "rag-chunk-v2.1"
EXPERT_REVIEW_VERSION = "final-expert-v5"

PUBLIC_TAIL = (
    "本规则适用于软件研发、电商运营、销售复盘、"
    "项目管理、市场营销等会议"
)

EXPECTED_COUNTS = {
    "POLICY_FIXED_24": 24,
    "RAG_DYNAMIC_32": 32,
    "RAG_SCENARIO_30": 30,
    "MERGE_MAP_20": 20,
    "ALIAS_TRIGGER_13": 13,
    "MEMORY_RULES_2": 2,
    "DOMAIN_EXTENSION_4": 4,
}


class ExportValidationError(RuntimeError):
    pass


def normalize_text(value: Any) -> str:
    if value is None:
        return ""

    return " ".join(str(value).strip().split())


def read_sheet(
    workbook: Any,
    sheet_name: str,
) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ExportValidationError(
            f"缺少工作表: {sheet_name}"
        )

    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))

    if not values:
        raise ExportValidationError(
            f"工作表为空: {sheet_name}"
        )

    headers = [
        normalize_text(value)
        for value in values[0]
    ]

    if not all(headers):
        raise ExportValidationError(
            f"工作表存在空表头: {sheet_name}"
        )

    rows: list[dict[str, Any]] = []

    for row_number, values_row in enumerate(
        values[1:],
        start=2,
    ):
        if all(value is None for value in values_row):
            continue

        row = {
            headers[index]: (
                values_row[index]
                if index < len(values_row)
                else None
            )
            for index in range(len(headers))
        }

        row["_row_number"] = row_number
        rows.append(row)

    return rows


def require_fields(
    rows: list[dict[str, Any]],
    sheet_name: str,
    fields: list[str],
) -> None:
    for row in rows:
        row_number = row["_row_number"]

        for field in fields:
            if not normalize_text(row.get(field)):
                raise ExportValidationError(
                    f"{sheet_name} 第 {row_number} 行"
                    f"字段为空: {field}"
                )


def make_rag_chunk(
    row: dict[str, Any],
) -> dict[str, Any]:
    chunk_id = normalize_text(row["chunk_id"])
    content = normalize_text(row["最终内容"])
    runtime_layer = normalize_text(row["运行位置"])
    layer_or_scenario = normalize_text(row["层级/场景"])
    dimension = normalize_text(row["维度"])
    knowledge_type = normalize_text(row["知识角色"])
    priority = normalize_text(row["优先级"]) or "medium"
    status = normalize_text(row["状态"]) or "approved"

    scenario = (
        layer_or_scenario
        if runtime_layer == "RAG_SCENARIO"
        else "all"
    )

    logical_layer = (
        layer_or_scenario
        if runtime_layer != "RAG_SCENARIO"
        else "scenario"
    )

    return {
        "id": chunk_id,
        "chunk_id": chunk_id,
        "title": normalize_text(row["title"]),
        "content": content,
        "metadata": {
            "runtime_layer": runtime_layer,
            "layer": logical_layer,
            "scenario": scenario,
            "dimension": dimension,
            "knowledge_type": knowledge_type,
            "priority": priority,
            "version": "2.1",
            "dataset_version": DATASET_VERSION,
            "chunk_schema_version": SCHEMA_VERSION,
            "expert_review_version": EXPERT_REVIEW_VERSION,
            "status": status,
        },
    }


def write_jsonl(
    path: Path,
    rows: list[dict[str, Any]],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    with path.open(
        "w",
        encoding="utf-8",
        newline="\n",
    ) as file:
        for row in rows:
            file.write(
                json.dumps(
                    row,
                    ensure_ascii=False,
                    sort_keys=False,
                )
            )
            file.write("\n")


def write_json(
    path: Path,
    value: Any,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    path.write_text(
        json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
            sort_keys=False,
        )
        + "\n",
        encoding="utf-8",
    )


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()

    with path.open("rb") as file:
        while True:
            block = file.read(1024 * 1024)

            if not block:
                break

            digest.update(block)

    return digest.hexdigest()


def validate_unique_ids(
    chunks: list[dict[str, Any]],
) -> None:
    ids = [chunk["id"] for chunk in chunks]
    duplicate_ids = sorted(
        chunk_id
        for chunk_id in set(ids)
        if ids.count(chunk_id) > 1
    )

    if duplicate_ids:
        raise ExportValidationError(
            "RAG 存在重复 ID: "
            + ", ".join(duplicate_ids)
        )


def validate_rag_chunks(
    chunks: list[dict[str, Any]],
) -> None:
    allowed_runtime_layers = {
        "POLICY_FIXED",
        "RAG_DYNAMIC",
        "RAG_SCENARIO",
    }

    allowed_dimensions = {
        "meeting_agenda",
        "meeting_summary",
        "key_conclusions",
        "action_items",
        "unresolved_issues",
        "risks_and_focus",
        "cross_dimension",
        "evidence",
    }

    allowed_roles = {
        "definition",
        "boundary_rule",
        "positive_example",
        "negative_example",
        "reasoning_pattern",
        "owner_resolution",
        "deadline_rule",
    }

    for chunk in chunks:
        chunk_id = chunk["id"]
        content = normalize_text(chunk["content"])
        metadata = chunk["metadata"]

        if not chunk_id:
            raise ExportValidationError(
                "发现空 chunk ID"
            )

        if not content:
            raise ExportValidationError(
                f"{chunk_id} 正文为空"
            )

        if PUBLIC_TAIL in content:
            raise ExportValidationError(
                f"{chunk_id} 仍包含公共尾句"
            )

        runtime_layer = metadata["runtime_layer"]
        dimension = metadata["dimension"]
        knowledge_type = metadata["knowledge_type"]

        if runtime_layer not in allowed_runtime_layers:
            raise ExportValidationError(
                f"{chunk_id} 非法 runtime_layer: "
                f"{runtime_layer}"
            )

        if dimension not in allowed_dimensions:
            raise ExportValidationError(
                f"{chunk_id} 非法 dimension: "
                f"{dimension}"
            )

        if knowledge_type not in allowed_roles:
            raise ExportValidationError(
                f"{chunk_id} 非法 knowledge_type: "
                f"{knowledge_type}"
            )

        if (
            runtime_layer == "RAG_SCENARIO"
            and metadata["scenario"] == "all"
        ):
            raise ExportValidationError(
                f"{chunk_id} 缺少场景 Metadata"
            )


def main() -> int:
    if not REVIEW_FILE.exists():
        raise FileNotFoundError(
            f"专家审查表不存在: {REVIEW_FILE}"
        )

    workbook = load_workbook(
        REVIEW_FILE,
        read_only=True,
        data_only=True,
    )

    policy_rows = read_sheet(
        workbook,
        "POLICY_FIXED_24",
    )

    dynamic_rows = read_sheet(
        workbook,
        "RAG_DYNAMIC_32",
    )

    scenario_rows = read_sheet(
        workbook,
        "RAG_SCENARIO_30",
    )

    merge_rows = read_sheet(
        workbook,
        "MERGE_MAP_20",
    )

    alias_rows = read_sheet(
        workbook,
        "ALIAS_TRIGGER_13",
    )

    memory_rows = read_sheet(
        workbook,
        "MEMORY_RULES_2",
    )

    domain_rows = read_sheet(
        workbook,
        "DOMAIN_EXTENSION_4",
    )

    actual_counts = {
        "POLICY_FIXED_24": len(policy_rows),
        "RAG_DYNAMIC_32": len(dynamic_rows),
        "RAG_SCENARIO_30": len(scenario_rows),
        "MERGE_MAP_20": len(merge_rows),
        "ALIAS_TRIGGER_13": len(alias_rows),
        "MEMORY_RULES_2": len(memory_rows),
        "DOMAIN_EXTENSION_4": len(domain_rows),
    }

    if actual_counts != EXPECTED_COUNTS:
        raise ExportValidationError(
            "工作表数量不匹配:\n"
            + json.dumps(
                {
                    "expected": EXPECTED_COUNTS,
                    "actual": actual_counts,
                },
                ensure_ascii=False,
                indent=2,
            )
        )

    rag_required_fields = [
        "chunk_id",
        "title",
        "最终内容",
        "运行位置",
        "层级/场景",
        "维度",
        "知识角色",
        "状态",
    ]

    require_fields(
        policy_rows,
        "POLICY_FIXED_24",
        rag_required_fields,
    )

    require_fields(
        dynamic_rows,
        "RAG_DYNAMIC_32",
        rag_required_fields,
    )

    require_fields(
        scenario_rows,
        "RAG_SCENARIO_30",
        rag_required_fields,
    )

    require_fields(
        merge_rows,
        "MERGE_MAP_20",
        [
            "source_id",
            "merge_target",
            "可吸收内容",
        ],
    )

    require_fields(
        alias_rows,
        "ALIAS_TRIGGER_13",
        [
            "source_id",
            "词典内容",
            "scope",
        ],
    )

    require_fields(
        memory_rows,
        "MEMORY_RULES_2",
        [
            "source_id",
            "规则内容",
        ],
    )

    require_fields(
        domain_rows,
        "DOMAIN_EXTENSION_4",
        [
            "source_id",
            "领域词",
        ],
    )

    policy_chunks = [
        make_rag_chunk(row)
        for row in policy_rows
    ]

    dynamic_chunks = [
        make_rag_chunk(row)
        for row in dynamic_rows
    ]

    scenario_chunks = [
        make_rag_chunk(row)
        for row in scenario_rows
    ]

    all_rag_chunks = (
        policy_chunks
        + dynamic_chunks
        + scenario_chunks
    )

    if len(all_rag_chunks) != 86:
        raise ExportValidationError(
            "独立 RAG 规则数量必须为 86，"
            f"实际为 {len(all_rag_chunks)}"
        )

    validate_unique_ids(all_rag_chunks)
    validate_rag_chunks(all_rag_chunks)

    final_ids = {
        chunk["id"]
        for chunk in all_rag_chunks
    }

    merge_map: list[dict[str, Any]] = []

    for row in merge_rows:
        source_id = normalize_text(row["source_id"])
        target_id = normalize_text(row["merge_target"])

        if target_id not in final_ids:
            raise ExportValidationError(
                f"merge_target 不存在于最终 RAG: "
                f"{source_id} -> {target_id}"
            )

        merge_map.append(
            {
                "source_id": source_id,
                "title": normalize_text(row["title"]),
                "merge_target": target_id,
                "absorbed_content": normalize_text(
                    row["可吸收内容"]
                ),
                "reason": normalize_text(
                    row["合并理由"]
                ),
            }
        )

    aliases = [
        {
            "source_id": normalize_text(row["source_id"]),
            "title": normalize_text(row["title"]),
            "content": normalize_text(row["词典内容"]),
            "scope": normalize_text(row["scope"]),
            "reason": normalize_text(row["迁移理由"]),
        }
        for row in alias_rows
    ]

    memory_rules = [
        {
            "id": normalize_text(row["source_id"]),
            "title": normalize_text(row["title"]),
            "content": normalize_text(row["规则内容"]),
            "metadata": {
                "runtime_layer": "MEMORY_RULES",
                "dataset_version": DATASET_VERSION,
                "expert_review_version": (
                    EXPERT_REVIEW_VERSION
                ),
                "reason": normalize_text(
                    row["迁移理由"]
                ),
            },
        }
        for row in memory_rows
    ]

    domain_extensions = [
        {
            "source_id": normalize_text(row["source_id"]),
            "title": normalize_text(row["title"]),
            "content": normalize_text(row["领域词"]),
            "status": "archived",
            "reason": normalize_text(row["归档理由"]),
        }
        for row in domain_rows
    ]

    write_jsonl(POLICY_FILE, policy_chunks)
    write_jsonl(DYNAMIC_FILE, dynamic_chunks)
    write_jsonl(SCENARIO_FILE, scenario_chunks)
    write_jsonl(COMBINED_FILE, all_rag_chunks)

    write_json(ALIAS_FILE, aliases)
    write_jsonl(MEMORY_FILE, memory_rules)
    write_json(DOMAIN_FILE, domain_extensions)
    write_json(MERGE_FILE, merge_map)

    output_files = [
        POLICY_FILE,
        DYNAMIC_FILE,
        SCENARIO_FILE,
        COMBINED_FILE,
        ALIAS_FILE,
        MEMORY_FILE,
        DOMAIN_FILE,
        MERGE_FILE,
    ]

    manifest = {
        "dataset_name": "MeetMind RAG v2.1",
        "dataset_version": DATASET_VERSION,
        "chunk_schema_version": SCHEMA_VERSION,
        "expert_review_version": EXPERT_REVIEW_VERSION,
        "generated_at": datetime.now(
            timezone.utc
        ).isoformat(),
        "source_review_file": str(
            REVIEW_FILE.relative_to(PROJECT_ROOT)
        ).replace("\\", "/"),
        "counts": {
            "policy_fixed": len(policy_chunks),
            "rag_dynamic": len(dynamic_chunks),
            "rag_scenario": len(scenario_chunks),
            "independent_rag_rules": len(
                all_rag_chunks
            ),
            "merge_map": len(merge_map),
            "alias_trigger": len(aliases),
            "memory_rules": len(memory_rules),
            "domain_extensions": len(
                domain_extensions
            ),
        },
        "files": {
            str(path.relative_to(PROJECT_ROOT)).replace(
                "\\",
                "/",
            ): {
                "sha256": sha256_file(path),
                "size_bytes": path.stat().st_size,
            }
            for path in output_files
        },
        "collection_recommendation": {
            "collection_name": (
                "meeting_analyst_rules_v2_1"
            ),
            "embedding_model": (
                "BAAI/bge-small-zh-v1.5"
            ),
            "import_file": str(
                COMBINED_FILE.relative_to(PROJECT_ROOT)
            ).replace("\\", "/"),
        },
        "retrieval_design": {
            "policy_fixed": (
                "固定选择，不依赖纯向量 Top-K"
            ),
            "rag_dynamic": (
                "根据会议正文动态召回"
            ),
            "rag_scenario": (
                "根据 meeting_type 和正文召回"
            ),
            "alias_trigger": (
                "用于查询扩展，不导入向量库"
            ),
            "memory_rules": (
                "进入 Agent Memory 独立链路"
            ),
        },
    }

    write_json(MANIFEST_FILE, manifest)

    print("RAG v2.1 export succeeded")
    print()
    print("Counts:")
    print(f"  POLICY_FIXED = {len(policy_chunks)}")
    print(f"  RAG_DYNAMIC  = {len(dynamic_chunks)}")
    print(f"  RAG_SCENARIO = {len(scenario_chunks)}")
    print(f"  RAG_TOTAL    = {len(all_rag_chunks)}")
    print(f"  MERGE_MAP    = {len(merge_map)}")
    print(f"  ALIAS        = {len(aliases)}")
    print(f"  MEMORY       = {len(memory_rules)}")
    print(
        f"  DOMAIN       = "
        f"{len(domain_extensions)}"
    )
    print()
    print(f"Manifest: {MANIFEST_FILE}")

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (
        ExportValidationError,
        FileNotFoundError,
        KeyError,
        ValueError,
    ) as error:
        print(
            f"RAG v2.1 export failed: {error}",
            file=sys.stderr,
        )
        raise SystemExit(1)